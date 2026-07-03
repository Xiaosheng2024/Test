# -*- coding: utf-8 -*-
"""
缺料汇总生成器 Qt5 版

功能：
1. 客户无需命令行，打开界面后选择“缺料报表详情导出*.xlsx”或任意符合格式的 Excel。
2. 点击“生成汇总表”，自动在导入文件同目录生成：汇总表_YYYYMMDD.xlsx。
3. 自动识别列名：零件号、零件名称、净库存、YYYY-MM-DD(班次一缺/班次二缺)。
4. 生成 3 个 Sheet：原始数据、白夜班、一整天。

依赖：
    pip install PyQt5 openpyxl

Mac 测试：
    python3 缺料汇总生成器_Qt5.py

Windows 运行源码：
    python 缺料汇总生成器_Qt5.py

Windows 打包 exe：
    pip install pyinstaller PyQt5 openpyxl
    pyinstaller --noconsole --onefile --name 缺料汇总生成器 缺料汇总生成器_Qt5.py
    生成的 exe 在 dist 目录下。
"""

import os
import re
import sys
import traceback
from pathlib import Path
from datetime import datetime

from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from PyQt5.QtCore import Qt, QThread, pyqtSignal
from PyQt5.QtWidgets import (
    QApplication,
    QWidget,
    QLabel,
    QPushButton,
    QLineEdit,
    QFileDialog,
    QMessageBox,
    QVBoxLayout,
    QHBoxLayout,
    QTextEdit,
    QProgressBar,
)


DATE_COL_RE = re.compile(r"^(\d{4}-\d{2}-\d{2})\(班次([一二])缺\)$")


def to_number(value):
    """把 Excel 单元格值转换成数字；空值返回 None。"""
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return value
    try:
        n = float(str(value).strip())
        return int(n) if n.is_integer() else n
    except Exception:
        return None


def is_zero_or_blank(values):
    """判断一组值是否全部为空或全部为 0。"""
    nums = [to_number(v) for v in values if to_number(v) is not None]
    return (not nums) or all(abs(float(v)) < 1e-12 for v in nums)


def read_source(input_file):
    wb = load_workbook(input_file, data_only=True)
    ws = wb.active

    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise ValueError("源表为空。")

    headers = [str(v).strip() if v is not None else "" for v in rows[0]]
    body = [list(r) for r in rows[1:]]

    try:
        part_idx = headers.index("零件号")
        name_idx = headers.index("零件名称")
        stock_idx = headers.index("净库存")
    except ValueError as e:
        raise ValueError("源表第一行必须包含表头：零件号、零件名称、净库存。") from e

    date_cols = []
    for idx, h in enumerate(headers):
        m = DATE_COL_RE.match(h)
        if m:
            date_cols.append({
                "idx": idx,
                "header": h,
                "date": m.group(1),
                "shift": m.group(2),
            })

    if not date_cols:
        raise ValueError("没有识别到日期班次列，例如：2026-07-07(班次一缺)。")

    max_base_idx = max(part_idx, name_idx, stock_idx)
    part_rows = [
        r for r in body
        if len(r) > max_base_idx
        and r[part_idx] not in (None, "")
        and r[name_idx] not in (None, "")
    ]

    if not part_rows:
        raise ValueError("没有找到包含零件号、零件名称的数据行。")

    # 跳过末尾全 0 或全空的无效日期班次列，避免最后几天无效列算出负数/异常。
    valid_cols = []
    for c in date_cols:
        values = [r[c["idx"]] if c["idx"] < len(r) else None for r in part_rows]
        if not is_zero_or_blank(values):
            valid_cols.append(c)

    if not valid_cols:
        raise ValueError("日期班次列存在，但全部为空或全部为 0，无法汇总。")

    return part_rows, part_idx, name_idx, stock_idx, valid_cols


def build_result(part_rows, part_idx, name_idx, stock_idx, valid_cols):
    base_headers = ["零件号", "零件名称", "净库存"]

    all_consumption_rows = []
    for r in part_rows:
        part_no = r[part_idx]
        part_name = r[name_idx]
        stock = to_number(r[stock_idx])
        base = [part_no, part_name, stock]

        prev_balance = stock
        consumption_values = []
        for c in valid_cols:
            cur_balance = to_number(r[c["idx"]]) if c["idx"] < len(r) else None

            if prev_balance is None or cur_balance is None:
                qty = None
            else:
                qty = prev_balance - cur_balance
                if isinstance(qty, float) and qty.is_integer():
                    qty = int(qty)

            consumption_values.append(qty)

            if cur_balance is not None:
                prev_balance = cur_balance

        all_consumption_rows.append(base + consumption_values)

    dates_in_order = []
    for c in valid_cols:
        if c["date"] not in dates_in_order:
            dates_in_order.append(c["date"])

    # 只保留有实际用量变化的日期；同一天班次一/班次二成对保留。
    keep_dates = []
    for date in dates_in_order:
        col_positions = [i for i, c in enumerate(valid_cols) if c["date"] == date]
        values = []
        for row in all_consumption_rows:
            for pos in col_positions:
                values.append(row[3 + pos])
        if not is_zero_or_blank(values):
            keep_dates.append(date)

    keep_positions = [i for i, c in enumerate(valid_cols) if c["date"] in keep_dates]

    raw_headers = base_headers + [valid_cols[i]["header"] for i in keep_positions]
    raw_rows = []
    for r in part_rows:
        row = [r[part_idx], r[name_idx], to_number(r[stock_idx])]
        for i in keep_positions:
            idx = valid_cols[i]["idx"]
            row.append(to_number(r[idx]) if idx < len(r) else None)
        raw_rows.append(row)

    shift_headers = base_headers + [valid_cols[i]["header"] for i in keep_positions]
    shift_rows = [
        row[:3] + [row[3 + i] for i in keep_positions]
        for row in all_consumption_rows
    ]

    daily_headers = base_headers + keep_dates
    daily_rows = []
    for row in all_consumption_rows:
        out = row[:3]
        for date in keep_dates:
            total = 0
            has_value = False
            for i, c in enumerate(valid_cols):
                if c["date"] == date:
                    v = row[3 + i]
                    if v is not None:
                        total += v
                        has_value = True
            out.append(total if has_value else None)
        daily_rows.append(out)

    return {
        "原始数据": (raw_headers, raw_rows),
        "白夜班": (shift_headers, shift_rows),
        "一整天": (daily_headers, daily_rows),
    }


def write_sheet(ws, headers, rows):
    ws.append(headers)
    for row in rows:
        ws.append(row)

    max_row = ws.max_row
    max_col = ws.max_column

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(bold=True, color="FFFFFF")
    thin = Side(style="thin", color="D9E2F3")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border

    for row in ws.iter_rows(min_row=2, max_row=max_row, max_col=max_col):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="center")
            if isinstance(cell.value, (int, float)):
                cell.number_format = "0"

    ws.freeze_panes = "D2"
    ws.auto_filter.ref = f"A1:{get_column_letter(max_col)}{max_row}"

    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 30
    ws.column_dimensions["C"].width = 10
    for col in range(4, max_col + 1):
        ws.column_dimensions[get_column_letter(col)].width = 15

    # 注意：这里不要创建 Excel Table 对象。
    # 部分 Mac/Windows Excel 对中文工作表 + 大量日期列的 Table 对象兼容性不好，
    # 打开时会提示“发现部分内容有问题”，修复后可能只剩 Sheet1。
    # 使用普通单元格 + 自动筛选 + 冻结窗格，兼容性更稳定。


def save_result(result, output_file):
    wb = Workbook()
    default = wb.active
    wb.remove(default)

    for sheet_name, (headers, rows) in result.items():
        ws = wb.create_sheet(sheet_name)
        write_sheet(ws, headers, rows)

    # 先保存到临时文件，保存成功后再替换成正式文件，避免中途异常留下坏文件。
    output_path = Path(output_file)
    temp_file = output_path.with_name(output_path.stem + "_tmp" + output_path.suffix)
    if temp_file.exists():
        temp_file.unlink()

    wb.save(temp_file)

    # 保存后立刻回读校验，确认 3 个 Sheet 都存在。
    check_wb = load_workbook(temp_file, read_only=True, data_only=True)
    expected_sheets = ["原始数据", "白夜班", "一整天"]
    if check_wb.sheetnames != expected_sheets:
        check_wb.close()
        temp_file.unlink(missing_ok=True)
        raise RuntimeError(f"生成文件校验失败，Sheet 不完整：{check_wb.sheetnames}")
    check_wb.close()

    if output_path.exists():
        output_path.unlink()
    temp_file.replace(output_path)


def make_output_path(input_file):
    input_path = Path(input_file)
    today = datetime.now().strftime("%Y%m%d")
    output = input_path.parent / f"汇总表_{today}.xlsx"

    # 如果当天已经生成过，自动追加 01、02，避免覆盖客户文件。
    if not output.exists():
        return str(output)

    for i in range(1, 100):
        candidate = input_path.parent / f"汇总表_{today}_{i:02d}.xlsx"
        if not candidate.exists():
            return str(candidate)

    raise RuntimeError("同一天生成文件过多，请清理目录后再试。")


def generate_summary(input_file):
    if not input_file:
        raise ValueError("请先选择导入文件。")
    if not os.path.exists(input_file):
        raise FileNotFoundError("导入文件不存在。")
    if not input_file.lower().endswith(".xlsx"):
        raise ValueError("请选择 .xlsx 格式的 Excel 文件。")

    output_file = make_output_path(input_file)
    part_rows, part_idx, name_idx, stock_idx, valid_cols = read_source(input_file)
    result = build_result(part_rows, part_idx, name_idx, stock_idx, valid_cols)
    save_result(result, output_file)

    return {
        "output_file": output_file,
        "row_count": len(part_rows),
        "shift_col_count": len(result["白夜班"][0]) - 3,
        "daily_col_count": len(result["一整天"][0]) - 3,
    }


class GenerateWorker(QThread):
    success = pyqtSignal(dict)
    failed = pyqtSignal(str)

    def __init__(self, input_file):
        super().__init__()
        self.input_file = input_file

    def run(self):
        try:
            result = generate_summary(self.input_file)
            self.success.emit(result)
        except Exception:
            self.failed.emit(traceback.format_exc())


class MainWindow(QWidget):
    def __init__(self):
        super().__init__()
        self.worker = None
        self.setWindowTitle("缺料汇总生成器")
        self.resize(720, 420)
        self.init_ui()

    def init_ui(self):
        title = QLabel("缺料汇总生成器")
        title.setAlignment(Qt.AlignCenter)
        title.setStyleSheet("font-size: 22px; font-weight: bold; padding: 10px;")

        tip = QLabel("选择客户导出的缺料报表 Excel，点击生成后，会在原文件同目录生成 汇总表_日期.xlsx。")
        tip.setWordWrap(True)
        tip.setStyleSheet("font-size: 13px; color: #555;")

        self.file_edit = QLineEdit()
        self.file_edit.setPlaceholderText("请选择 .xlsx 文件")
        self.file_edit.setReadOnly(True)

        browse_btn = QPushButton("选择导入文件")
        browse_btn.clicked.connect(self.choose_file)

        file_layout = QHBoxLayout()
        file_layout.addWidget(self.file_edit)
        file_layout.addWidget(browse_btn)

        self.generate_btn = QPushButton("生成汇总表")
        self.generate_btn.setMinimumHeight(44)
        self.generate_btn.setStyleSheet("font-size: 16px; font-weight: bold;")
        self.generate_btn.clicked.connect(self.start_generate)

        self.progress = QProgressBar()
        self.progress.setRange(0, 1)
        self.progress.setValue(0)
        self.progress.setTextVisible(False)

        self.log = QTextEdit()
        self.log.setReadOnly(True)
        self.log.setPlaceholderText("处理日志会显示在这里。")

        layout = QVBoxLayout()
        layout.addWidget(title)
        layout.addWidget(tip)
        layout.addLayout(file_layout)
        layout.addWidget(self.generate_btn)
        layout.addWidget(self.progress)
        layout.addWidget(self.log)
        self.setLayout(layout)

    def choose_file(self):
        file_path, _ = QFileDialog.getOpenFileName(
            self,
            "选择缺料报表 Excel",
            "",
            "Excel 文件 (*.xlsx)"
        )
        if file_path:
            self.file_edit.setText(file_path)
            self.log.append(f"已选择文件：{file_path}")

    def start_generate(self):
        input_file = self.file_edit.text().strip()
        if not input_file:
            QMessageBox.warning(self, "提示", "请先选择导入文件。")
            return

        self.generate_btn.setEnabled(False)
        self.progress.setRange(0, 0)  # 忙碌状态
        self.log.append("开始生成汇总表，请稍候……")

        self.worker = GenerateWorker(input_file)
        self.worker.success.connect(self.on_success)
        self.worker.failed.connect(self.on_failed)
        self.worker.start()

    def on_success(self, result):
        self.progress.setRange(0, 1)
        self.progress.setValue(1)
        self.generate_btn.setEnabled(True)

        msg = (
            "生成成功！\n\n"
            f"输出文件：{result['output_file']}\n"
            f"零件行数：{result['row_count']}\n"
            f"白夜班列数：{result['shift_col_count']}\n"
            f"一整天日期数：{result['daily_col_count']}"
        )
        self.log.append(msg.replace("\n", "\n"))
        QMessageBox.information(self, "完成", msg)

    def on_failed(self, error_text):
        self.progress.setRange(0, 1)
        self.progress.setValue(0)
        self.generate_btn.setEnabled(True)

        self.log.append("生成失败：")
        self.log.append(error_text)
        QMessageBox.critical(self, "生成失败", "生成失败，请检查导入表格式。详细错误已显示在日志区域。")


def main():
    QApplication.setAttribute(Qt.AA_EnableHighDpiScaling, True)
    app = QApplication(sys.argv)
    win = MainWindow()
    win.show()
    sys.exit(app.exec_())


if __name__ == "__main__":
    main()
